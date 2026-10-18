"""Ground truth from data/open_labels.xlsx: file | таймстемп | нарушение / без нарушения | тип нарушения (rule_id)."""
import os, re

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "data", "open_labels.xlsx")
TYPE_COL = "тип нарушения"

# rule_id per violation clip (plan table). X-4 and X-47 are the same file labelled at two moments.
GT_TYPES = {
    "X-1": "cash_return_violation", "X-2": "anonymous_cash_acceptance_violation",
    "Х-4": "currency_intermix_violation", "Х-11": "left_key_violation", "Х-13": "uncounted_safe_deposit",
    "Х-15": "left_collection_bag_violation", "Х-17": "anonymous_incassator_violation",
    "Х-18": "interrupted_transaction", "Х-20": "open_safe_violation", "Х-22": "left_key_violation",
    "Х-24": "left_commodities_violation", "Х-26": "left_money_violation", "Х-28": "left_laptop_violation",
    "Х-30": "open_door_violation_intruder", "Х-32": "open_door_violation", "Х-33": "open_door_violation_intruder",
    "Х-34": "interrupted_transaction", "Х-38": "interrupted_transaction",
    "Х-40": "personal_belongings_violation", "Х-44": "currency_intermix_violation",
    "Х-46": "beverage_violation", "Х-47": "interrupted_transaction",
}


def parse_timestamp(s):
    """'22 секунда' -> 22.0, '1минута 8 секунда' -> 68.0, None/'' -> None."""
    if s is None or str(s).strip() == "":
        return None
    s = str(s).lower()
    m = re.search(r"(\d+(?:[.,]\d+)?)\s*мин", s)
    sec = re.search(r"(\d+(?:[.,]\d+)?)\s*сек", s)
    if not m and not sec:
        num = re.search(r"\d+(?:[.,]\d+)?", s)
        return float(num.group().replace(",", ".")) if num else None
    val = lambda x: float(x.group(1).replace(",", ".")) if x else 0.0
    return val(m) * 60 + val(sec)


def load_gt(path=XLSX):
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or "").strip() for h in rows[0]]
    ti = head.index(TYPE_COL) if TYPE_COL in head else None
    out = {}
    for r in rows[1:]:
        if not r[0]: continue
        v = str(r[2] or "").strip().lower() == "нарушение"
        out[str(r[0]).strip()] = {"verdict": "violation" if v else "no_violation",
                                  "timestamp_sec": parse_timestamp(r[1]),
                                  "type": (str(r[ti]).strip() if ti is not None and r[ti] else None)}
    return out


def add_type_column(path=XLSX):
    """Write the тип нарушения column (rule_id slugs; empty for clean rows) into the Excel file."""
    import openpyxl
    wb = openpyxl.load_workbook(path); ws = wb.active
    head = [c.value for c in ws[1]]
    col = head.index(TYPE_COL) + 1 if TYPE_COL in head else ws.max_column + 1
    ws.cell(1, col, TYPE_COL)
    for row in range(2, ws.max_row + 1):
        name, verdict = ws.cell(row, 1).value, str(ws.cell(row, 3).value or "")
        ws.cell(row, col, GT_TYPES.get(str(name).strip()) if verdict.strip() == "нарушение" else None)
    wb.save(path)


def compare_ground_truth(gt, alarms):
    """gt: {"verdict", "timestamp_sec", "type"}; alarms: list of dicts/objects with rule_id and timecode_sec."""
    get = lambda a, k: a[k] if isinstance(a, dict) else getattr(a, k)
    pred_types = sorted({get(a, "rule_id") for a in alarms})
    verdict = "violation" if alarms else "no_violation"
    exp = gt.get("type")
    type_match = (exp in pred_types) if exp else not pred_types
    err = None
    if gt.get("timestamp_sec") is not None and alarms:
        match = [a for a in alarms if get(a, "rule_id") == exp]
        a = min(match or alarms[:1], key=lambda a: abs(get(a, "timecode_sec") - gt["timestamp_sec"]))
        err = round(abs(get(a, "timecode_sec") - gt["timestamp_sec"]), 1)
    return {"verdict_match": verdict == gt["verdict"], "type_match": type_match,
            "timestamp_abs_error_sec": err, "predicted_types": pred_types, "predicted_verdict": verdict}


if __name__ == "__main__":
    add_type_column()
    for k, v in load_gt().items(): print(k, v)
